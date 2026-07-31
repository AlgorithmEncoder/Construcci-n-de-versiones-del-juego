"""
XLSX reader.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook


def read_xlsx(path: Path) -> list[list[str]]:

    workbook = load_workbook(
        path,
        data_only=True
    )

    sheet = workbook.active

    rows = []

    for row in sheet.iter_rows(values_only=True):

        rows.append(

            [

                "" if value is None else str(value)

                for value in row

            ]

        )

    return rows