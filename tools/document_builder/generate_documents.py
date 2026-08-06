"""
===========================================================
Dream Office - Document Builder
Version 1.0

Reads:
    files.json
    computer_files.json

Generates:
    PDF
    DOCX
    XLSX
    PPTX
    TXT

Automatically skips unchanged documents using a manifest.

===========================================================
"""

from pathlib import Path
import json
import hashlib
import shutil
import traceback

from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    PageBreak
)

from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.units import cm

from docx import Document
from docx.shared import Pt

from openpyxl import Workbook
from openpyxl.styles import Font

from pptx import Presentation
from pptx.util import Inches, Pt as PPTPt


# ==========================================================
# CONFIGURATION
# ==========================================================

# -----------------------------
# CHANGE THESE PATHS
# -----------------------------

ROOT = Path(__file__).parent.parent.parent

MEMORY_FOLDER = ROOT / "data" / "memories" / "memory_01"

FILES_JSON = MEMORY_FOLDER / "files.json"

COMPUTER_FILES_JSON = MEMORY_FOLDER / "computer_files.json"

OUTPUT_FOLDER = MEMORY_FOLDER / "assets" / "documents"

MANIFEST_FILE = Path(__file__).parent / ".manifest.json"

# -----------------------------

COMPANY_NAME = "Northbridge Solutions"

DEFAULT_AUTHOR = "Document Builder"

DEFAULT_CLASSIFICATION = "Internal"

DEFAULT_REVISION = "1.0"

SKIP_UNCHANGED = True

VERBOSE = True


# ==========================================================
# STYLES
# ==========================================================

PDF_STYLES = getSampleStyleSheet()

PDF_STYLES["Title"].alignment = TA_CENTER

PDF_STYLES["BodyText"].leading = 18

PDF_STYLES["BodyText"].spaceAfter = 8


# ==========================================================
# DOCUMENT BUILDER
# ==========================================================

class DocumentBuilder:

    def __init__(self):

        self.files = {}

        self.computer_files = {}

        self.manifest = {}

        self.created = 0

        self.updated = 0

        self.skipped = 0

        self.errors = 0

    # =====================================================

    def info(self, text):

        if VERBOSE:

            print(text)

    # =====================================================

    def ensure_folder(self, folder):

        Path(folder).mkdir(
            parents=True,
            exist_ok=True
        )

    # =====================================================

    def load_json(self, path):

        with open(path, "r", encoding="utf8") as f:

            return json.load(f)

    # =====================================================

    def save_json(self, path, data):

        with open(path, "w", encoding="utf8") as f:

            json.dump(
                data,
                f,
                indent=4,
                ensure_ascii=False
            )

    # =====================================================

    def canonical_json(self, obj):

        return json.dumps(
            obj,
            sort_keys=True,
            separators=(",", ":"),
            ensure_ascii=False
        )

    # =====================================================

    def calculate_hash(self, obj):

        return hashlib.sha256(
            self.canonical_json(obj).encode("utf8")
        ).hexdigest()

    # =====================================================

    def load(self):

        self.files = self.load_json(
            FILES_JSON
        )

        self.computer_files = self.load_json(
            COMPUTER_FILES_JSON
        )

        if MANIFEST_FILE.exists():

            self.manifest = self.load_json(
                MANIFEST_FILE
            )

        else:

            self.manifest = {}

    # =====================================================

    def save_manifest(self):

        self.save_json(
            MANIFEST_FILE,
            self.manifest
        )

    # =====================================================

    def output_path(self, document_id):

        relative = self.computer_files[
            document_id
        ]["file"]

        return OUTPUT_FOLDER / relative

    # =====================================================

    def extension(self, document_id):

        return self.output_path(
            document_id
        ).suffix.lower()

    # =====================================================

    def document_hash(self, document_id):

        return self.calculate_hash(

            self.files[document_id]

        )

    # =====================================================

    def changed(self, document_id):

        new_hash = self.document_hash(
            document_id
        )

        old = self.manifest.get(
            document_id
        )

        if old is None:

            return True

        if old["hash"] != new_hash:

            return True

        if not self.output_path(
            document_id
        ).exists():

            return True

        return False
    
        # =====================================================
    # STATUS
    # =====================================================

    def created_file(self, path):

        self.created += 1

        print(f"[CREATED ] {path}")

    # =====================================================

    def updated_file(self, path):

        self.updated += 1

        print(f"[UPDATED ] {path}")

    # =====================================================

    def skipped_file(self, path):

        self.skipped += 1

        print(f"[SKIPPED ] {path}")

    # =====================================================

    def error_file(self, path):

        self.errors += 1

        print(f"[ERROR   ] {path}")

    # =====================================================
    # HELPERS
    # =====================================================

    def document_data(self, document_id):

        return self.files[document_id]

    # =====================================================

    def template(self, document_id):

        return self.document_data(document_id).get(
            "template",
            "default"
        )

    # =====================================================

    def file_type(self, document_id):

        ext = self.extension(document_id)

        mapping = {

            ".pdf": "pdf",

            ".docx": "docx",

            ".xlsx": "xlsx",

            ".pptx": "pptx",

            ".txt": "txt"

        }

        return mapping.get(ext)

    # =====================================================

    def ensure_defaults(self, document):

        document.setdefault(
            "author",
            DEFAULT_AUTHOR
        )

        document.setdefault(
            "classification",
            DEFAULT_CLASSIFICATION
        )

        document.setdefault(
            "revision",
            DEFAULT_REVISION
        )

        document.setdefault(
            "content",
            []
        )

    # =====================================================

    def prepare_output(self, document_id):

        output = self.output_path(
            document_id
        )

        self.ensure_folder(
            output.parent
        )

        return output

    # =====================================================

    def update_manifest_entry(self, document_id):

        self.manifest[document_id] = {

            "hash": self.document_hash(
                document_id
            ),

            "file": self.computer_files[
                document_id
            ]["file"]

        }

    # =====================================================
    # GENERATION
    # =====================================================

    def generate_document(self, document_id):

        document = self.document_data(
            document_id
        )

        self.ensure_defaults(
            document
        )

        output = self.prepare_output(
            document_id
        )

        if SKIP_UNCHANGED:

            if not self.changed(
                document_id
            ):

                self.skipped_file(
                    output
                )

                return

        filetype = self.file_type(
            document_id
        )

        if filetype == "pdf":

            self.generate_pdf(
                document_id,
                output
            )

        elif filetype == "docx":

            self.generate_docx(
                document_id,
                output
            )

        elif filetype == "xlsx":

            self.generate_xlsx(
                document_id,
                output
            )

        elif filetype == "pptx":

            self.generate_pptx(
                document_id,
                output
            )

        elif filetype == "txt":

            self.generate_txt(
                document_id,
                output
            )

        else:

            raise Exception(

                f"Unsupported type: {filetype}"

            )

        self.update_manifest_entry(
            document_id
        )

    # =====================================================

    def generate_all(self):

        total = len(
            self.files
        )

        print()

        print("=" * 60)

        print(
            "Generating documents..."
        )

        print("=" * 60)

        print()

        for index, document_id in enumerate(

            sorted(self.files),

            start=1

        ):

            print(

                f"[{index}/{total}] {document_id}"

            )

            try:

                self.generate_document(
                    document_id
                )

            except Exception:

                self.error_file(
                    document_id
                )

                traceback.print_exc()

                print()

        self.save_manifest()

        print()

        print("=" * 60)

        print("Generation finished")

        print("=" * 60)

        print()

        print(f"Created : {self.created}")

        print(f"Updated : {self.updated}")

        print(f"Skipped : {self.skipped}")

        print(f"Errors  : {self.errors}")

        print()
    
        # =====================================================
    # CONTENT HELPERS
    # =====================================================

    def blocks(self, document_id):

        return self.document_data(document_id).get(
            "content",
            []
        )

    # =====================================================

    def write_heading(self, file, text):

        file.write(text.upper())
        file.write("\n")
        file.write("=" * len(text))
        file.write("\n\n")

    # =====================================================

    def write_paragraph(self, file, text):

        file.write(text)
        file.write("\n\n")

    # =====================================================

    def write_list(self, file, items):

        for item in items:

            file.write(f"• {item}\n")

        file.write("\n")

    # =====================================================

    def write_table(self, file, headers, rows):

        widths = []

        for h in headers:
            widths.append(len(str(h)))

        for row in rows:

            for i, value in enumerate(row):

                widths[i] = max(
                    widths[i],
                    len(str(value))
                )

        def line(values):

            pieces = []

            for i, value in enumerate(values):

                pieces.append(
                    str(value).ljust(widths[i])
                )

            return " | ".join(pieces)

        file.write(line(headers))
        file.write("\n")

        file.write(
            "-+-".join(
                "-" * w
                for w in widths
            )
        )

        file.write("\n")

        for row in rows:

            file.write(line(row))

            file.write("\n")

        file.write("\n")

    # =====================================================
    # TXT GENERATOR
    # =====================================================

    def generate_txt(
        self,
        document_id,
        output
    ):

        document = self.document_data(
            document_id
        )

        existed = output.exists()

        with open(
            output,
            "w",
            encoding="utf8"
        ) as file:

            if document.get("title"):

                self.write_heading(
                    file,
                    document["title"]
                )

            for block in self.blocks(
                document_id
            ):

                block_type = block["type"]

                if block_type == "heading":

                    self.write_heading(
                        file,
                        block["text"]
                    )

                elif block_type == "paragraph":

                    self.write_paragraph(
                        file,
                        block["text"]
                    )

                elif block_type == "bullet_list":

                    self.write_list(
                        file,
                        block["items"]
                    )

                elif block_type == "table":

                    self.write_table(
                        file,
                        block["headers"],
                        block["rows"]
                    )

                elif block_type == "page_break":

                    file.write("\n")
                    file.write("=" * 60)
                    file.write("\n\n")

                else:

                    file.write(
                        f"[Unsupported block: {block_type}]"
                    )

                    file.write("\n\n")

        if existed:

            self.updated_file(output)

        else:

            self.created_file(output)
    
        # =====================================================
    # PDF HELPERS
    # =====================================================

    def pdf_title(self, story, text):

        story.append(
            Paragraph(
                text,
                PDF_STYLES["Title"]
            )
        )

        story.append(
            Spacer(1, 0.5 * cm)
        )

    # =====================================================

    def pdf_heading(self, story, text):

        story.append(
            Paragraph(
                text,
                PDF_STYLES["Heading2"]
            )
        )

        story.append(
            Spacer(1, 0.2 * cm)
        )

    # =====================================================

    def pdf_paragraph(self, story, text):

        story.append(
            Paragraph(
                text,
                PDF_STYLES["BodyText"]
            )
        )

    # =====================================================

    def pdf_list(self, story, items):

        for item in items:

            story.append(

                Paragraph(

                    f"• {item}",

                    PDF_STYLES["BodyText"]

                )

            )

        story.append(
            Spacer(1, 0.2 * cm)
        )

    # =====================================================

    def pdf_table(self, story, headers, rows):

        data = [headers]

        data.extend(rows)

        table = Table(data)

        table.setStyle(

            TableStyle(

                [

                    (
                        "BACKGROUND",
                        (0, 0),
                        (-1, 0),
                        colors.lightgrey
                    ),

                    (
                        "GRID",
                        (0, 0),
                        (-1, -1),
                        0.5,
                        colors.black
                    ),

                    (
                        "FONTNAME",
                        (0, 0),
                        (-1, 0),
                        "Helvetica-Bold"
                    ),

                    (
                        "BOTTOMPADDING",
                        (0, 0),
                        (-1, 0),
                        8
                    ),

                    (
                        "TOPPADDING",
                        (0, 1),
                        (-1, -1),
                        4
                    ),

                    (
                        "BOTTOMPADDING",
                        (0, 1),
                        (-1, -1),
                        4
                    )

                ]

            )

        )

        story.append(table)

        story.append(
            Spacer(1, 0.4 * cm)
        )

    # =====================================================
    # PDF GENERATOR
    # =====================================================

    def generate_pdf(
        self,
        document_id,
        output
    ):

        document = self.document_data(
            document_id
        )

        existed = output.exists()

        pdf = SimpleDocTemplate(
            str(output)
        )

        story = []

        if document.get("title"):

            self.pdf_title(
                story,
                document["title"]
            )

        info = (
            f"<b>Classification:</b> "
            f"{document['classification']}<br/>"

            f"<b>Revision:</b> "
            f"{document['revision']}<br/>"

            f"<b>Author:</b> "
            f"{document['author']}"
        )

        story.append(

            Paragraph(

                info,

                PDF_STYLES["BodyText"]

            )

        )

        story.append(
            Spacer(1, 0.5 * cm)
        )

        for block in self.blocks(
            document_id
        ):

            block_type = block["type"]

            if block_type == "heading":

                self.pdf_heading(
                    story,
                    block["text"]
                )

            elif block_type == "paragraph":

                self.pdf_paragraph(
                    story,
                    block["text"]
                )

            elif block_type == "bullet_list":

                self.pdf_list(
                    story,
                    block["items"]
                )

            elif block_type == "table":

                self.pdf_table(
                    story,
                    block["headers"],
                    block["rows"]
                )

            elif block_type == "page_break":

                story.append(
                    PageBreak()
                )

            else:

                story.append(

                    Paragraph(

                        f"[Unsupported block: {block_type}]",

                        PDF_STYLES["BodyText"]

                    )

                )

        pdf.build(story)

        if existed:

            self.updated_file(output)

        else:

            self.created_file(output)

        # =====================================================
    # DOCX HELPERS
    # =====================================================

    def docx_title(self, doc, text):

        p = doc.add_heading(level=0)
        p.alignment = 1
        p.add_run(text)

    # =====================================================

    def docx_heading(self, doc, text):

        doc.add_heading(text, level=1)

    # =====================================================

    def docx_paragraph(self, doc, text):

        p = doc.add_paragraph()

        run = p.add_run(text)

        run.font.size = Pt(11)

    # =====================================================

    def docx_list(self, doc, items):

        for item in items:

            doc.add_paragraph(
                item,
                style="List Bullet"
            )

    # =====================================================

    def docx_table(self, doc, headers, rows):

        table = doc.add_table(
            rows=1,
            cols=len(headers)
        )

        table.style = "Table Grid"

        header_cells = table.rows[0].cells

        for i, value in enumerate(headers):

            header_cells[i].text = str(value)

        for row in rows:

            cells = table.add_row().cells

            for i, value in enumerate(row):

                cells[i].text = str(value)

    # =====================================================
    # DOCX GENERATOR
    # =====================================================

    def generate_docx(
        self,
        document_id,
        output
    ):

        document = self.document_data(
            document_id
        )

        existed = output.exists()

        doc = Document()

        # -------------------------------------------------

        if document.get("title"):

            self.docx_title(
                doc,
                document["title"]
            )

        info = doc.add_paragraph()

        info.add_run("Classification: ").bold = True
        info.add_run(document["classification"])

        info.add_run("\nRevision: ").bold = True
        info.add_run(document["revision"])

        info.add_run("\nAuthor: ").bold = True
        info.add_run(document["author"])

        doc.add_paragraph()

        # -------------------------------------------------

        for block in self.blocks(
            document_id
        ):

            block_type = block["type"]

            if block_type == "heading":

                self.docx_heading(
                    doc,
                    block["text"]
                )

            elif block_type == "paragraph":

                self.docx_paragraph(
                    doc,
                    block["text"]
                )

            elif block_type == "bullet_list":

                self.docx_list(
                    doc,
                    block["items"]
                )

            elif block_type == "table":

                self.docx_table(
                    doc,
                    block["headers"],
                    block["rows"]
                )

            elif block_type == "page_break":

                doc.add_page_break()

            else:

                self.docx_paragraph(

                    doc,

                    f"[Unsupported block: {block_type}]"

                )

        doc.save(output)

        if existed:

            self.updated_file(output)

        else:

            self.created_file(output)
    
        # =====================================================
    # XLSX HELPERS
    # =====================================================

    def xlsx_heading(self, ws, row, text):

        cell = ws.cell(row=row, column=1)

        cell.value = text

        cell.font = Font(
            bold=True,
            size=14
        )

        return row + 2

    # =====================================================

    def xlsx_paragraph(self, ws, row, text):

        ws.cell(
            row=row,
            column=1
        ).value = text

        return row + 2

    # =====================================================

    def xlsx_list(self, ws, row, items):

        for item in items:

            ws.cell(
                row=row,
                column=1
            ).value = f"• {item}"

            row += 1

        return row + 1

    # =====================================================

    def xlsx_table(self, ws, row, headers, rows):

        col = 1

        for header in headers:

            cell = ws.cell(
                row=row,
                column=col
            )

            cell.value = header

            cell.font = Font(
                bold=True
            )

            col += 1

        row += 1

        for values in rows:

            col = 1

            for value in values:

                ws.cell(
                    row=row,
                    column=col
                ).value = value

                col += 1

            row += 1

        return row + 1

    # =====================================================
    # XLSX GENERATOR
    # =====================================================

    def generate_xlsx(
        self,
        document_id,
        output
    ):

        document = self.document_data(
            document_id
        )

        existed = output.exists()

        workbook = Workbook()

        sheet = workbook.active

        sheet.title = "Document"

        row = 1

        # -------------------------------------------------

        if document.get("title"):

            sheet["A1"] = document["title"]

            sheet["A1"].font = Font(
                bold=True,
                size=18
            )

            row = 3

        sheet.cell(
            row=row,
            column=1
        ).value = "Classification"

        sheet.cell(
            row=row,
            column=2
        ).value = document["classification"]

        row += 1

        sheet.cell(
            row=row,
            column=1
        ).value = "Revision"

        sheet.cell(
            row=row,
            column=2
        ).value = document["revision"]

        row += 1

        sheet.cell(
            row=row,
            column=1
        ).value = "Author"

        sheet.cell(
            row=row,
            column=2
        ).value = document["author"]

        row += 3

        # -------------------------------------------------

        for block in self.blocks(
            document_id
        ):

            block_type = block["type"]

            if block_type == "heading":

                row = self.xlsx_heading(
                    sheet,
                    row,
                    block["text"]
                )

            elif block_type == "paragraph":

                row = self.xlsx_paragraph(
                    sheet,
                    row,
                    block["text"]
                )

            elif block_type == "bullet_list":

                row = self.xlsx_list(
                    sheet,
                    row,
                    block["items"]
                )

            elif block_type == "table":

                row = self.xlsx_table(
                    sheet,
                    row,
                    block["headers"],
                    block["rows"]
                )

            elif block_type == "page_break":

                row += 3

            else:

                sheet.cell(
                    row=row,
                    column=1
                ).value = f"[Unsupported block: {block_type}]"

                row += 2

        workbook.save(output)

        if existed:

            self.updated_file(output)

        else:

            self.created_file(output)
    
    # =====================================================
    # PPTX HELPERS
    # =====================================================

    def generate_business_presentation(
        self,
        presentation,
        document
    ):

        layout = presentation.slide_layouts[1]  # Title + Content

        for slide_data in document.get("content", []):

            if slide_data.get("type") != "slide":
                continue

            slide = presentation.slides.add_slide(layout)

            slide.shapes.title.text = slide_data.get(
                "title",
                ""
            )

            body = slide.placeholders[1].text_frame
            body.clear()

            first = True

            for line in slide_data.get("body", []):

                if first:

                    p = body.paragraphs[0]
                    first = False

                else:

                    p = body.add_paragraph()

                p.text = line
                p.level = 0
                p.font.size = PPTPt(22)

    def ppt_new_slide(self, presentation, title=""):

        layout = presentation.slide_layouts[1]

        slide = presentation.slides.add_slide(layout)

        title_box = slide.shapes.title

        if title_box is not None:

            title_box.text = title

        return slide

    # =====================================================

    def ppt_add_text(self, slide, text, level=0):

        left = Inches(0.7)
        top = Inches(1.0)
        width = Inches(8.2)
        height = Inches(5.0)

        textbox = slide.shapes.add_textbox(
            left,
            top,
            width,
            height
        )

        frame = textbox.text_frame

        paragraph = frame.paragraphs[0]

        paragraph.text = text

        paragraph.level = level

        paragraph.font.size = PPTPt(20)

    # =====================================================

    def ppt_add_list(self, slide, items):

        left = Inches(0.8)
        top = Inches(1.2)
        width = Inches(8)
        height = Inches(5)

        textbox = slide.shapes.add_textbox(
            left,
            top,
            width,
            height
        )

        frame = textbox.text_frame

        first = True

        for item in items:

            if first:

                p = frame.paragraphs[0]

                first = False

            else:

                p = frame.add_paragraph()

            p.text = item

            p.level = 0

            p.font.size = PPTPt(20)

    # =====================================================

    def ppt_add_table(self, slide, headers, rows):

        rows_count = len(rows) + 1

        cols_count = len(headers)

        table = slide.shapes.add_table(

            rows_count,

            cols_count,

            Inches(0.5),

            Inches(1.2),

            Inches(8.3),

            Inches(0.8 + rows_count * 0.35)

        ).table

        for col, value in enumerate(headers):

            table.cell(0, col).text = str(value)

        for r, values in enumerate(rows, start=1):

            for c, value in enumerate(values):

                table.cell(r, c).text = str(value)

    # =====================================================
    # PPTX GENERATOR
    # =====================================================

    def generate_pptx(
        self,
        document_id,
        output
    ):

        document = self.document_data(document_id)

        existed = output.exists()

        presentation = Presentation()

        # Eliminar diapositiva inicial
        while len(presentation.slides):

            rId = presentation.slides._sldIdLst[0].rId

            presentation.part.drop_rel(rId)

            del presentation.slides._sldIdLst[0]

        # ============================================
        # PRESENTACIONES
        # ============================================

        if document.get("template") == "presentation":

            self.generate_business_presentation(
                presentation,
                document
            )

        # ============================================
        # DOCUMENTOS NORMALES
        # ============================================

        else:

            slide = self.ppt_new_slide(
                presentation,
                document.get("title", "")
            )

            for block in self.blocks(document_id):

                block_type = block["type"]

                if block_type == "heading":

                    slide = self.ppt_new_slide(
                        presentation,
                        block["text"]
                    )

                elif block_type == "paragraph":

                    self.ppt_add_text(
                        slide,
                        block["text"]
                    )

                elif block_type == "bullet_list":

                    self.ppt_add_list(
                        slide,
                        block["items"]
                    )

                elif block_type == "table":

                    self.ppt_add_table(
                        slide,
                        block["headers"],
                        block["rows"]
                    )

                elif block_type == "page_break":

                    slide = self.ppt_new_slide(
                        presentation
                    )

                else:

                    self.ppt_add_text(
                        slide,
                        f"[Unsupported block: {block_type}]"
                    )

        presentation.save(output)

        if existed:

            self.updated_file(output)

        else:

            self.created_file(output)
    
    # =====================================================
    # BINARY FILES
    # =====================================================

    def generate_binary(
        self,
        document_id,
        output
    ):

        document = self.document_data(
            document_id
        )

        source = OUTPUT_FOLDER / document["source"]

        if not source.exists():

            raise FileNotFoundError(source)

        existed = output.exists()

        shutil.copy2(
            source,
            output
        )

        if existed:

            self.updated_file(output)

        else:

            self.created_file(output)

    # =====================================================
    # GENERATOR DISPATCH
    # =====================================================

    def generate_document(self, document_id):

        document = self.document_data(document_id)

        self.ensure_defaults(document)

        output = self.prepare_output(document_id)

        if SKIP_UNCHANGED:

            if not self.changed(document_id):

                self.skipped_file(output)

                return

        filetype = self.file_type(document_id)

        if filetype == "pdf":

            self.generate_pdf(
                document_id,
                output
            )

        elif filetype == "docx":

            self.generate_docx(
                document_id,
                output
            )

        elif filetype == "xlsx":

            self.generate_xlsx(
                document_id,
                output
            )

        elif filetype == "pptx":

            self.generate_pptx(
                document_id,
                output
            )

        elif filetype == "txt":

            self.generate_txt(
                document_id,
                output
            )

        elif self.extension(document_id) in (

            ".png",

            ".jpg",

            ".jpeg",

            ".zip"

        ):

            self.generate_binary(
                document_id,
                output
            )

        else:

            raise Exception(

                f"Unsupported extension: {self.extension(document_id)}"

            )

        self.update_manifest_entry(
            document_id
        )

    # =====================================================
    # VALIDATION
    # =====================================================

    def validate(self):

        print()

        print("Validating...")

        print()

        errors = False

        for document_id in self.files:

            if document_id not in self.computer_files:

                print(

                    f"[ERROR] Missing in computer_files.json: {document_id}"

                )

                errors = True

        for document_id in self.computer_files:

            if document_id not in self.files:

                print(

                    f"[ERROR] Missing in files.json: {document_id}"

                )

                errors = True

        if not errors:

            print("Validation successful.")

        print()
    
    # =====================================================
        # UTILITIES
        # =====================================================
    
        def exists(self, document_id):
    
            return self.output_path(
                document_id
            ).exists()
    
        # =====================================================
    
        def document_summary(self, document_id):
    
            document = self.document_data(
                document_id
            )
    
            return {
    
                "title": document.get(
                    "title",
                    document_id
                ),
    
                "template": document.get(
                    "template",
                    "default"
                ),
    
                "blocks": len(
                    document.get(
                        "content",
                        []
                    )
                ),
    
                "extension": self.extension(
                    document_id
                )
    
            }
    
        # =====================================================
    
        def print_summary(self):
    
            print()
            print("=" * 60)
            print("DOCUMENT SUMMARY")
            print("=" * 60)
    
            for document_id in sorted(self.files):
    
                info = self.document_summary(
                    document_id
                )
    
                print(
    
                    f"{document_id:30}"
    
                    f"{info['extension']:>8}   "
    
                    f"{info['blocks']:3} blocks"
    
                )
    
            print("=" * 60)
            print()
    
        # =====================================================
    
        def verify_output(self):
    
            missing = []
    
            for document_id in self.files:
    
                if not self.exists(document_id):
    
                    missing.append(
                        document_id
                    )
    
            if missing:
    
                print()
    
                print("Missing generated files:")
    
                for document_id in missing:
    
                    print(
                        " -",
                        document_id
                    )
    
                print()
    
                return False
    
            return True


# ==========================================================
# MAIN
# ==========================================================

def print_banner():

    print("=" * 60)
    print("Dream Office - Document Builder")
    print("Version 1.0")
    print("=" * 60)
    print()


def check_configuration():

    if not FILES_JSON.exists():
        raise FileNotFoundError(
            f"Missing: {FILES_JSON}"
        )

    if not COMPUTER_FILES_JSON.exists():
        raise FileNotFoundError(
            f"Missing: {COMPUTER_FILES_JSON}"
        )

    OUTPUT_FOLDER.mkdir(
        parents=True,
        exist_ok=True
    )


def run():

    print_banner()

    builder = DocumentBuilder()

    check_configuration()

    builder.load()

    builder.validate()

    builder.generate_all()

    builder.print_summary()

    if builder.verify_output():

        print("All documents generated successfully.")

    else:

        print("Some documents could not be generated.")

    return builder


# ==========================================================
# ENTRY POINT
# ==========================================================

if __name__ == "__main__":

    try:

        run()

        print()
        print("=" * 60)
        print("Finished successfully.")
        print("=" * 60)

    except KeyboardInterrupt:

        print()
        print("Generation cancelled by user.")

    except Exception:

        print()
        print("=" * 60)
        print("Fatal error")
        print("=" * 60)
        print()

        traceback.print_exc()

        print()

        raise SystemExit(1)
    
    